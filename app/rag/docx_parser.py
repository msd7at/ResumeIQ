from docx import Document
from pathlib import Path


def parse_docx(file_path: str) -> str:
    """
    Extract all text from a DOCX file.
    Reads paragraphs and table cells. Returns a single cleaned string.
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"DOCX not found: {file_path}")
    if path.suffix.lower() != ".docx":
        raise ValueError(f"Expected a .docx file, got: {path.suffix}")

    doc = Document(file_path)
    parts = []

    # paragraphs — covers headings, bullets, normal text
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)

    # tables — some resumes put skills/experience in table cells
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    parts.append(text)

    full_text = "\n".join(parts)
    return _clean_text(full_text)


def _clean_text(text: str) -> str:
    """Collapse multiple blank lines and remove non-printable characters."""
    lines = text.splitlines()
    cleaned = []
    prev_blank = False

    for line in lines:
        line = line.strip()
        line = "".join(ch for ch in line if ch.isprintable())
        is_blank = len(line) == 0

        if is_blank and prev_blank:
            continue
        cleaned.append(line)
        prev_blank = is_blank

    return "\n".join(cleaned).strip()


if __name__ == "__main__":
    import sys
    if len(sys.argv) < 2:
        print("Usage: python docx_parser.py <path_to_resume.docx>")
        sys.exit(1)

    text = parse_docx(sys.argv[1])
    print("── Extracted Text ──────────────────────")
    print(text[:2000])
    print(f"\n── Total characters extracted: {len(text)}")
